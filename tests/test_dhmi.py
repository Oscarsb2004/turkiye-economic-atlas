"""
The airport readers, and what they refuse.

DHMİ publishes traffic against a NAME and no code, in a workbook that carries
two years side by side and a third block of percentage changes that mentions
both of them. Most of what can go wrong here is reading the wrong column or
quietly losing an airport, and neither would look wrong on a map.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from atlas.core import registry as R
from atlas.readers import dhmi, ourairports

# Three real airports and one real total row, which is all the shape needs.
AIRPORTS = [("İstanbul", "LTFM"), ("Ankara Esenboğa", "LTAC"), ("Trabzon", "LTCG")]


def _workbook(year: int, previous: int, *, rows=None, total=None, sheets=None,
              headers=None) -> bytes:
    """A workbook shaped the way DHMİ shapes its own, with the values under our control."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    values = rows if rows is not None else {name: (10, 1, 11) for name, _ in AIRPORTS}
    for printed in (sheets if sheets is not None else dhmi.SHEETS):
        sheet = book.create_sheet(printed)
        sheet.cell(1, 1, f"{printed} TRAFİĞİ")
        sheet.cell(2, 1, "Havalimanları ")
        head = headers if headers is not None else (f"{previous} ARALIK SONU\n", f"{year} ARALIK SONU\n",
                                                   f" {year}/{previous} (%)")
        for block, text in enumerate(head):
            sheet.cell(2, 2 + block * 3, text)
            for offset, label in enumerate(("İç Hat", "Dış Hat", "Toplam")):
                sheet.cell(3, 2 + block * 3 + offset, label)
        row = 4
        for name in values:
            sheet.cell(row, 1, name)
            for offset, value in enumerate(values[name]):
                sheet.cell(row, 2 + offset, 0)                 # the previous year
                sheet.cell(row, 5 + offset, value)             # the year wanted
            row += 1
        summed = total if total is not None else tuple(
            sum(v[i] for v in values.values()) for i in range(3))
        sheet.cell(row, 1, "DHMİ TOPLAMI")
        sheet.cell(row + 1, 1, "TÜRKİYE GENELİ")
        for offset, value in enumerate(summed):
            sheet.cell(row + 1, 5 + offset, value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_the_year_is_found_by_its_header_not_its_position():
    traffic = dhmi.annual(_workbook(2025, 2024), 2025)
    assert sorted(traffic.by_icao) == ["LTAC", "LTCG", "LTFM"]
    assert traffic.by_icao["LTFM"]["passengers_domestic"] == 10
    assert traffic.by_icao["LTFM"]["passengers_total"] == 11
    assert traffic.published_total["passengers_total"] == 33

    # The previous year is in the same file, and reading it reads different
    # columns rather than the same ones.
    assert dhmi.annual(_workbook(2025, 2024), 2024).by_icao["LTFM"]["passengers_total"] == 0


def test_the_percentage_block_is_not_mistaken_for_a_year():
    """
    Its header is " 2025/2024 (%)", so a search for either year finds it. Taking
    it would publish a percentage change as if it were a count.
    """
    body = _workbook(2025, 2024)
    traffic = dhmi.annual(body, 2025)
    # The block that was read is the one with counts in it, not the one with
    # percentages: 10 + 1 = 11 is the published total, not a ratio.
    assert traffic.by_icao["LTAC"]["passengers_domestic"] == 10


def test_a_year_the_workbook_does_not_carry_is_refused():
    with pytest.raises(dhmi.DhmiError, match="no column block for 2019"):
        dhmi.annual(_workbook(2025, 2024), 2019)


def test_an_airport_the_crosswalk_does_not_know_is_refused():
    """A new airport, a rename or a respelling has to be a deliberate line."""
    rows = {"İstanbul": (1, 1, 2), "Yeni Havalimanı": (1, 1, 2)}
    with pytest.raises(dhmi.DhmiError, match="registry/airports.yaml"):
        dhmi.annual(_workbook(2025, 2024, rows=rows, total=(2, 2, 4)), 2025)


def test_a_missing_national_row_is_refused():
    """Without it nothing holds the parts against the publisher's own sum."""
    body = _workbook(2025, 2024)
    book = openpyxl.load_workbook(io.BytesIO(body))
    for sheet in book:
        for row in range(1, sheet.max_row + 1):
            if str(sheet.cell(row, 1).value or "").strip() == "TÜRKİYE GENELİ":
                sheet.cell(row, 1, "")
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(dhmi.DhmiError, match="nothing checks the parts"):
        dhmi.annual(buffer.getvalue(), 2025)


def test_a_missing_sheet_is_refused():
    with pytest.raises(dhmi.DhmiError, match="no sheet"):
        dhmi.annual(_workbook(2025, 2024, sheets=["YOLCU"]), 2025)


def test_the_crosswalk_holds_every_name_exactly_once():
    """Two airports claiming one of DHMİ's names would make the join ambiguous."""
    names = R.airport_by_dhmi_name()
    assert names["İstanbul"] == "LTFM", "the airport that opened in 2019, not Atatürk"
    assert names["İstanbul Atatürk"] == "LTBA"
    # A rename and a respelling, both pointing at one airport.
    assert names["Erzincan"] == names["Erzincan Yıldırım Akbulut"] == "LTCD"
    assert names["Şanlıurfa GAP"] == names["Şanlıurfa Gap"] == "LTCS"


# ── OurAirports ────────────────────────────────────────────────────────────────

HEADER = ("id,ident,type,name,latitude_deg,longitude_deg,elevation_ft,continent,iso_country,"
          "iso_region,municipality,scheduled_service,icao_code,iata_code,gps_code,local_code,"
          "home_link,wikipedia_link,keywords\n")


def _csv(*rows: str) -> bytes:
    return (HEADER + "".join(rows)).encode("utf-8")


ISTANBUL = ("1,LTFM,large_airport,İstanbul Airport,41.274874,28.732136,325,EU,TR,TR-34,"
            "Arnavutköy,yes,LTFM,IST,,,,,\n")


def test_the_province_comes_from_the_publishers_own_region():
    found = ourairports.turkish(_csv(ISTANBUL), {"LTFM"})
    assert found["LTFM"].plaka == 34
    assert found["LTFM"].iata == "IST"
    assert found["LTFM"].lon == 28.732136


def test_an_airport_the_file_does_not_carry_is_refused():
    with pytest.raises(ourairports.OurAirportsError, match="LTXX"):
        ourairports.turkish(_csv(ISTANBUL), {"LTFM", "LTXX"})


def test_a_region_that_is_not_a_plaka_code_is_refused():
    """TR-U-A and the like: the province cannot be read from it, so it is not guessed."""
    odd = ISTANBUL.replace("TR-34", "TR-U-A")
    with pytest.raises(ourairports.OurAirportsError, match="iso_region"):
        ourairports.turkish(_csv(odd), {"LTFM"})
