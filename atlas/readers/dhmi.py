"""
atlas.readers.dhmi — airport traffic, from DHMİ's comparative workbooks.

    https://www.dhmi.gov.tr/Sayfalar/Istatistikler.aspx

WHAT ONE WORKBOOK HOLDS

Five sheets — all aircraft, passengers, commercial aircraft, freight, cargo —
each laid out for a reader:

    row 1   the sheet's title, in words
    row 2   TWO period headers, merged over three columns each, and a third
            block of percentage changes: "2024 ARALIK SONU", "2025 ARALIK SONU",
            " 2025/2024 (%)"
    row 3   İç Hat, Dış Hat, Toplam under each block
    rows 4+ one airport each, by NAME and no code
    then    DHMİ TOPLAMI, TÜRKİYE GENELİ, and four direct-transit rows
    last    two footnotes

THE FIGURES ARE CUMULATIVE, so December is the year. The file for one year also
republishes the previous year beside it, which is why the year wanted is found
by READING THE HEADERS: taking a column position would silently read the wrong
year the first time DHMİ adds a column, and the percentage block carries BOTH
years in its own header, so a naive search for "2025" finds it too.

TWO TOTALS, AND ONLY ONE OF THEM ADDS UP

DHMİ TOPLAMI excludes the airports it marks with (*) — İstanbul, Sabiha Gökçen,
Çukurova and four smaller ones are operated by others — so it is NOT the sum of
the rows above it. TÜRKİYE GENELİ is. That is the one reproduced here for the
declared check to hold the parts against.

NAMES, WHICH IS WHY THERE IS A CROSSWALK

The first column is a name, and DHMİ renames ("Erzincan" became "Erzincan
Yıldırım Akbulut"), respells ("Şanlıurfa GAP" / "Şanlıurfa Gap") and moves the
space in its own footnote marker ("İstanbul(*)" became "İstanbul (*)"). The
marker is stripped; everything else is looked up in registry/airports.yaml, and
a name that is not there stops the run.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass

import openpyxl

from atlas.core import registry as R

log = logging.getLogger(__name__)

#: The sheets, and what this project calls each measure. Matched on the sheet
#: name with its whitespace stripped: 2020's workbook names one sheet "YÜK "
#: and 2025's names it "YÜK".
SHEETS = {
    "TÜM UÇAK": "aircraft",
    "TİCARİ UÇAK": "commercial_aircraft",
    "YOLCU": "passengers",
    "YÜK": "freight_tonnes",
    "KARGO": "cargo_tonnes",
}

#: The three columns under each period header, and what this project calls them.
SLICES = {"İç Hat": "domestic", "Dış Hat": "international", "Toplam": "total"}

#: The row that closes the airports, and the one that totals them all. DHMİ
#: TOPLAMI leaves out the airports marked (*); TÜRKİYE GENELİ does not.
END_OF_AIRPORTS = "DHMİ TOPLAMI"
NATIONAL_ROW = "TÜRKİYE GENELİ"


class DhmiError(ValueError):
    """A DHMİ workbook is not shaped the way this reader was written for."""


@dataclass(frozen=True, slots=True)
class Traffic:
    """One year of traffic, as published."""

    year: str
    #: ICAO -> "<measure>_<slice>" -> value, exactly as the sheets print it.
    by_icao: dict[str, dict[str, float]]
    #: The same keys, from DHMİ's own TÜRKİYE GENELİ row.
    published_total: dict[str, float]
    #: ICAO -> the name DHMİ printed for it in this year's tables.
    names: dict[str, str]


def _year_columns(sheet, year: int) -> dict[str, int]:
    """Slice -> column, for the block whose header names `year` and nothing else."""
    start: int | None = None
    for col in range(1, sheet.max_column + 1):
        header = str(sheet.cell(2, col).value or "")
        years = set(re.findall(r"(?:19|20)\d\d", header))
        # The percentage block's header carries both years, which is exactly
        # the trap: it would match a search for either one.
        if "%" in header or len(years) != 1:
            continue
        if years == {str(year)}:
            start = col
            break
    if start is None:
        headers = [str(sheet.cell(2, c).value or "").strip() for c in range(1, sheet.max_column + 1)]
        raise DhmiError(f"{sheet.title}: no column block for {year}; the headers are {[h for h in headers if h]}")

    found: dict[str, int] = {}
    for offset in range(3):
        label = str(sheet.cell(3, start + offset).value or "").strip()
        if label not in SLICES:
            raise DhmiError(f"{sheet.title}: column {start + offset} under {year} is {label!r}, not one of {sorted(SLICES)}")
        found[SLICES[label]] = start + offset
    return found


def _number(value) -> float:
    """A published figure as it reads: an integer where it is one, a tonnage otherwise."""
    if value is None or value == "":
        return 0
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise DhmiError(f"{value!r} is not a number")
    return int(value) if float(value).is_integer() else float(value)


def annual(body: bytes, year: int) -> Traffic:
    """Every airport's traffic in `year`, and DHMİ's own total of it."""
    book = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
    sheets = {str(name).strip(): name for name in book.sheetnames}
    missing = sorted(set(SHEETS) - set(sheets))
    if missing:
        raise DhmiError(f"the workbook has no sheet(s) named {missing}; it has {book.sheetnames}")

    known = R.airport_by_dhmi_name()
    by_icao: dict[str, dict[str, float]] = {}
    published_total: dict[str, float] = {}
    names: dict[str, str] = {}

    for printed, measure in SHEETS.items():
        sheet = book[sheets[printed]]
        columns = _year_columns(sheet, year)
        seen: set[str] = set()
        national = None

        for row in range(4, sheet.max_row + 1):
            label = str(sheet.cell(row, 1).value or "").strip()
            if not label:
                continue
            if label.startswith(END_OF_AIRPORTS):
                # Everything below is a total or a footnote; find the one that
                # actually totals the rows above.
                for below in range(row, sheet.max_row + 1):
                    if str(sheet.cell(below, 1).value or "").strip() == NATIONAL_ROW:
                        national = {f"{measure}_{name}": _number(sheet.cell(below, col).value)
                                    for name, col in columns.items()}
                        break
                break

            # DHMİ marks privately operated airports with a trailing (*), and
            # moved the space in front of it in 2022. The marker is about who
            # operates the airport, not what it is called.
            name = re.sub(r"\s*\(\*\)\s*$", "", label)
            icao = known.get(name)
            if icao is None:
                raise DhmiError(
                    f"{year} {printed}: {name!r} is not in registry/airports.yaml. "
                    f"A new airport, a rename or a respelling is a line in that file"
                )
            if icao in seen:
                raise DhmiError(f"{year} {printed}: two rows for {icao} ({name!r})")
            seen.add(icao)
            names.setdefault(icao, name)
            figures = by_icao.setdefault(icao, {})
            for slice_name, col in columns.items():
                figures[f"{measure}_{slice_name}"] = _number(sheet.cell(row, col).value)

        if national is None:
            raise DhmiError(f"{year} {printed}: no {NATIONAL_ROW!r} row, so nothing checks the parts")
        published_total.update(national)

        if seen != set(by_icao):
            short = sorted(set(by_icao) - seen)
            raise DhmiError(f"{year} {printed}: no row for {short}, which other sheets carry")

    log.info("%s: %d airports, %d figures each", year, len(by_icao), len(next(iter(by_icao.values()))))
    return Traffic(year=str(year), by_icao=by_icao, published_total=published_total, names=names)
